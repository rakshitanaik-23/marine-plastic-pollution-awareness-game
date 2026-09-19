import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank

    # Color Palette
    BG_DARK = RGBColor(5, 16, 38)        # #051026 Deep Ocean Navy
    CARD_BG = RGBColor(12, 29, 61)       # #0c1d3d Card Navy
    CARD_BORDER = RGBColor(24, 52, 99)   # #183463 Border
    CYAN = RGBColor(0, 210, 255)         # #00d2ff Neon Ocean Cyan
    GREEN = RGBColor(16, 185, 129)       # #10b981 Emerald
    CORAL = RGBColor(255, 77, 109)       # #ff4d6d Coral Warning
    GOLD = RGBColor(255, 209, 102)       # #ffd166 Sand Gold
    WHITE = RGBColor(255, 255, 255)      # White
    GRAY = RGBColor(160, 180, 210)       # Light Slate
    TOXIC = RGBColor(57, 255, 20)        # #39ff14 Radioactive Green

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="SDG 14: LIFE BELOW WATER"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # ==========================================================
    # SLIDE 1: Title Slide
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Decorative Card in center
    s1_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    s1_card.fill.solid()
    s1_card.fill.fore_color.rgb = CARD_BG
    s1_card.line.color.rgb = CYAN
    s1_card.line.width = Pt(2)

    # Text Frame
    tf1 = s1_card.text_frame
    tf1.word_wrap = True
    tf1.margin_top = Inches(0.4)

    p1_badge = tf1.paragraphs[0]
    p1_badge.alignment = PP_ALIGN.CENTER
    p1_badge.text = "GLOBAL HACKATHON | UN SDG 14: LIFE BELOW WATER"
    p1_badge.font.size = Pt(13)
    p1_badge.font.bold = True
    p1_badge.font.color.rgb = GREEN

    p1_title = tf1.add_paragraph()
    p1_title.alignment = PP_ALIGN.CENTER
    p1_title.text = "OCEAN GUARDIAN 🐢🌊"
    p1_title.font.size = Pt(46)
    p1_title.font.bold = True
    p1_title.font.color.rgb = CYAN

    p1_tag = tf1.add_paragraph()
    p1_tag.alignment = PP_ALIGN.CENTER
    p1_tag.text = "“Save the Ocean, One Turtle at a Time.”"
    p1_tag.font.size = Pt(20)
    p1_tag.font.italic = True
    p1_tag.font.color.rgb = GOLD

    p1_desc = tf1.add_paragraph()
    p1_desc.alignment = PP_ALIGN.CENTER
    p1_desc.text = "\nA Modern 2D Marine Plastic Pollution Awareness & Educational Arcade Game\nBuilt with HTML5 Canvas, Web Audio API & Java Spring Boot REST Backend"
    p1_desc.font.size = Pt(14)
    p1_desc.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 2: Problem Statement & SDG 14
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "The Global Crisis: Marine Plastic Pollution")

    cols = [
        ("⚠️ The Scale", "14+ Million Tons", "Over 14 million metric tons of plastic enter our oceans annually, making up 80% of all marine debris from coastlines to deep ocean trenches.", CORAL),
        ("🐢 Endangered Wildlife", "100,000+ Deaths", "Sea turtles mistake floating plastic bags for jellyfish. Over 100,000 marine mammals and 1M seabirds die each year from plastic ingestion and entanglement.", GOLD),
        ("⏱️ The Longevity", "450+ Years", "A single PET plastic bottle takes up to 450 years to break down into toxic microplastics that bioaccumulate throughout the global food chain.", CYAN)
    ]

    for idx, (head, metric, desc, color) in enumerate(cols):
        left = Inches(0.8 + idx * 4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.4)

        p0 = tf.paragraphs[0]
        p0.text = head
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = color

        p1 = tf.add_paragraph()
        p1.text = metric
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 3: The Ocean Guardian Solution
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "The Solution: Gamified Environmental Education")

    pillars = [
        ("🎮 Experiential Empathy", "Puts players directly in the role of an endangered sea turtle fighting for survival in an increasingly hostile ocean environment."),
        ("🔍 Plastic vs. Food Differentiation", "Teaches players how easily marine creatures mistake transparent bags and wrappers for essential food like jellyfish and algae."),
        ("🌊 Dynamic Environmental Feedback", "Visually degrades the ocean ecosystem in real time—transitioning from pristine blue reefs to toxic industrial sludge."),
        ("💡 Measurable Awareness Impact", "Features actionable SDG 14 environmental facts and single-use plastic reduction principles embedded into every gameplay transition.")
    ]

    for idx, (title, desc) in enumerate(pillars):
        x = Inches(0.8 + (idx % 2) * 5.9)
        y = Inches(1.8 + (idx // 2) * 2.5)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = CYAN

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(13)
        p1.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 4: Core Gameplay Mechanics
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Game Mechanics & Hazard Interactions")

    mechanics = [
        ("🐟 Healthy Food Items", "+10 Points", "Jellyfish, Coral Fish & Kelp Pods", "Collecting nutritious food triggers joyful audio chimes and particle sparkle bursts. Helps hit early level completion thresholds.", GREEN),
        ("🧴 Plastic Debris", "-1 Life (Heart)", "Bottles, Bags, Wrappers, Cans", "Colliding with plastic triggers a hollow impact thud, screen shake, red flash, and 1.4s invulnerability blinking.", CORAL),
        ("☣️ Toxic Waste Barrels", "-2 Lives (Hearts)", "Industrial Radioactive Drums", "Introduced in Levels 4 & 5. Emits a pulsing neon green glow, toxic drip trail, and catastrophic double damage.", TOXIC),
        ("❤️ 3-Heart Health System", "Max 3 Lives", "HUD Display: ❤️❤️❤️", "Health reaching 0 triggers Game Over with environmental awareness reflection. Surviving all 5 levels crowns the Ocean Guardian.", GOLD)
    ]

    for idx, (title, stat, items, desc, col) in enumerate(mechanics):
        left = Inches(0.8 + idx * 2.95)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.3)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = col

        p1 = tf.add_paragraph()
        p1.text = stat
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p_items = tf.add_paragraph()
        p_items.text = items
        p_items.font.size = Pt(11)
        p_items.font.bold = True
        p_items.font.color.rgb = CYAN

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 5: 5 Progressive Levels (5-Minute Gameplay)
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "5-Level Progressive Duration System (5 Minutes Max)")

    levels_info = [
        ("Level 1", "Clean Waters", "0–100", "Easy", "Few plastic bottles, slow drift, abundant food, bright blue ocean.", GREEN),
        ("Level 2", "Rising Pollution", "101–250", "Medium", "Bottles & bags, faster movement, slightly murky waters.", CYAN),
        ("Level 3", "Plastic Storm", "251–450", "Hard", "Wrappers, cans, 6-pack rings, surging plastic debris waves.", GOLD),
        ("Level 4", "Toxic Ocean", "451–700", "Very Hard", "Toxic waste barrels (-2 ❤️), dark violet sludge, dead corals.", CORAL),
        ("Level 5", "Ocean Guardian", "701+", "Extreme", "Max plastic density, toxic sludge, final survival challenge.", RGBColor(168, 85, 247))
    ]

    for idx, (lvl, name, score, diff, desc, col) in enumerate(levels_info):
        left = Inches(0.8 + idx * 2.37)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(2.25), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = lvl
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = col

        p1 = tf.add_paragraph()
        p1.text = name
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = f"🎯 {score} pts\n⏱️ 60s Timer\n⚡ {diff}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = GOLD

        p3 = tf.add_paragraph()
        p3.text = "\n" + desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 6: Visual Progression & Audio Architecture
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Visual Progression & Procedural Audio Engine")

    # Left: Visual Progression Card
    left_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = CARD_BORDER

    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.3)

    pl0 = tf_l.paragraphs[0]
    pl0.text = "🎨 Dynamic Atmospheric Shaders"
    pl0.font.size = Pt(18)
    pl0.font.bold = True
    pl0.font.color.rgb = CYAN

    pl1 = tf_l.add_paragraph()
    pl1.text = "\n• Level 1: Crystal-clear turquoise water with sunbeams filtering down and vibrant blooming corals.\n• Level 2: Slightly desaturated teal waters with suspended dust particles.\n• Level 3: Gray-green polluted waters, microplastic particulate haze, and bleached corals.\n• Levels 4-5: Toxic violet-sludge abyss with zero sunlight, dead reefs, and glowing radioactive barrels.\n• 100% Procedural Vector Rendering: Crisp high-DPI graphics on any screen with zero image asset loading lag."
    pl1.font.size = Pt(12.5)
    pl1.font.color.rgb = GRAY

    # Right: Audio Engine Card
    right_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = CARD_BG
    right_card.line.color.rgb = CARD_BORDER

    tf_r = right_card.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    pr0 = tf_r.paragraphs[0]
    pr0.text = "🔊 Procedural Web Audio Engine"
    pr0.font.size = Pt(18)
    pr0.font.bold = True
    pr0.font.color.rgb = GREEN

    pr1 = tf_r.add_paragraph()
    pr1.text = "\n• Zero External Audio Dependencies: Built entirely on the native HTML5 Web Audio API.\n• Melodic Chimes: 4-note ascending C-major arpeggio on food collection.\n• Plastic Collision Thud: Sub-bass square wave combined with bandpassed white noise crack.\n• Toxic Hazard Alarm: Dual-sawtooth dissonance with sizzling high-pass noise.\n• Dynamic Fanfares: Countdown beeps (3-2-1-GO!), Level Complete chords, and Victory Fanfare."
    pr1.font.size = Pt(12.5)
    pr1.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 7: Technical Stack & Architecture
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "System Architecture & Technology Stack")

    tech_boxes = [
        ("Frontend Client", "HTML5, CSS3, JavaScript ES6+", "HTML5 Canvas 2D API for 60 FPS rendering, CSS glassmorphism, responsive 16:9 aspect ratio locking.", CYAN),
        ("Audio Synthesizer", "Web Audio API (Oscillators & Noise)", "Procedural real-time sound generation without requiring external MP3/WAV files. Works completely offline.", GREEN),
        ("Java REST Backend", "Java 11+ & Spring Boot REST API", "Lightweight HTTP server providing /api/scores, /api/sdg-facts, and static asset serving on port 8080.", GOLD),
        ("Hybrid Persistence", "REST API + LocalStorage Fallback", "Dual-mode persistence ensures scores and leaderboard always save even when running purely offline in browser.", CORAL)
    ]

    for idx, (title, tech, desc, col) in enumerate(tech_boxes):
        x = Inches(0.8 + (idx % 2) * 5.9)
        y = Inches(1.8 + (idx // 2) * 2.5)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = col

        p1 = tf.add_paragraph()
        p1.text = tech
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = GRAY

    # ==========================================================
    # SLIDE 8: Hackathon Impact & Conclusion
    # ==========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Impact, Demo & Conclusion")

    # Center Hero Card
    s8_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.6), Inches(10.333), Inches(5.1))
    s8_card.fill.solid()
    s8_card.fill.fore_color.rgb = CARD_BG
    s8_card.line.color.rgb = GREEN
    s8_card.line.width = Pt(2)

    tf8 = s8_card.text_frame
    tf8.word_wrap = True
    tf8.margin_left = Inches(0.5)
    tf8.margin_right = Inches(0.5)
    tf8.margin_top = Inches(0.4)

    p8_title = tf8.paragraphs[0]
    p8_title.alignment = PP_ALIGN.CENTER
    p8_title.text = "🌊 BECOME AN OCEAN GUARDIAN 🐢"
    p8_title.font.size = Pt(28)
    p8_title.font.bold = True
    p8_title.font.color.rgb = GOLD

    p8_quote = tf8.add_paragraph()
    p8_quote.alignment = PP_ALIGN.CENTER
    p8_quote.text = "“Every piece of plastic we prevent from reaching the ocean can make a difference.”"
    p8_quote.font.size = Pt(16)
    p8_quote.font.italic = True
    p8_quote.font.color.rgb = WHITE

    p8_body = tf8.add_paragraph()
    p8_body.text = """
✅ Fully Playable 2D Web Game: 5 progressive levels, 60s countdown timer, 3-2-1-GO transitions.
✅ Dual Architecture: Runs standalone in any browser (index.html) or with Java Server (http://localhost:8080).
✅ Zero Dependencies: Procedural vector graphics and Web Audio API synthesizer.
✅ SDG 14 Impact: Educational problem, solution, impact cards and interactive high-score leaderboard.
"""
    p8_body.font.size = Pt(13)
    p8_body.font.color.rgb = GRAY

    p8_footer = tf8.add_paragraph()
    p8_footer.alignment = PP_ALIGN.CENTER
    p8_footer.text = "Thank You! | Live Demo: http://localhost:8080"
    p8_footer.font.size = Pt(18)
    p8_footer.font.bold = True
    p8_footer.font.color.rgb = CYAN

    # Save presentation
    output_path = "D:\\microohackathon\\Ocean_Guardian_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == "__main__":
    create_deck()
